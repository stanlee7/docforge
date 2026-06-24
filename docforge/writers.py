"""생성된 문서 텍스트를 행별 파일로 기록 (docx / txt / md)."""
import os
import re

from . import config, excel_io

_MD_BOLD = re.compile(r"\*\*([^*]+)\*\*|__([^_]+)__")
_MD_HEAD = re.compile(r"^\s{0,3}#{1,6}\s*", re.MULTILINE)
_MD_BULLET = re.compile(r"^\s{0,3}[*+]\s+", re.MULTILINE)


def strip_markdown(text: str) -> str:
    """docx/txt용: 모델이 남긴 마크다운 표식(**굵게**, # 제목, * 글머리)을 평문화한다."""
    out = _MD_BOLD.sub(lambda m: m.group(1) or m.group(2), text or "")
    out = _MD_HEAD.sub("", out)
    out = _MD_BULLET.sub("- ", out)
    return out


def build_out_path(out_dir: str, filename_pattern: str, row: dict, fmt: str) -> str:
    """파일명 패턴을 행 데이터로 채우고 확장자를 붙여 충돌 없는 경로 반환."""
    fmt = fmt if fmt in config.OUTPUT_FORMATS else config.DEFAULT_FORMAT
    name = excel_io.fill_filename(filename_pattern, row)
    path = os.path.join(out_dir, f"{name}.{fmt}")
    return excel_io.unique_path(path)


def _write_docx(text: str, out_path: str, title: str = "") -> None:
    from docx import Document  # 지연 import: txt/md만 쓸 땐 python-docx 불필요
    from docx.shared import Pt

    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "맑은 고딕"
    style.font.size = Pt(11)
    if title:
        doc.add_heading(title, level=1)
    text = strip_markdown(text)
    for para in text.split("\n\n"):
        para = para.strip()
        if para:
            doc.add_paragraph(para)
    doc.save(out_path)


def _write_text(text: str, out_path: str, title: str = "", md: bool = False) -> None:
    body = text if md else strip_markdown(text)
    if title:
        head = f"# {title}\n\n" if md else f"{title}\n\n"
        body = head + body
    with open(out_path, "w", encoding="utf-8-sig") as f:
        f.write(body)


def write_document(text: str, out_path: str, fmt: str, title: str = "") -> str:
    """text를 fmt 형식으로 out_path에 기록하고 경로를 반환한다."""
    os.makedirs(os.path.dirname(os.path.abspath(out_path)), exist_ok=True)
    if fmt == "docx":
        _write_docx(text, out_path, title)
    elif fmt == "md":
        _write_text(text, out_path, title, md=True)
    else:  # txt (기본 폴백)
        _write_text(text, out_path, title, md=False)
    return out_path
