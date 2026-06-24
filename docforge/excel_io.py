"""엑셀 읽기 + {{컬럼}} 템플릿 머지 (mail-batch-tool engine.py에서 발췌, SMTP 제외)."""
import os
import re

import openpyxl

PLACEHOLDER = re.compile(r"\{\{([^{}]+)\}\}")
_ILLEGAL_FILENAME = re.compile(r'[\\/:*?"<>|\r\n\t]')


def sanitize_filename(name: str) -> str:
    name = _ILLEGAL_FILENAME.sub("_", str(name)).strip().rstrip(".")
    return name[:120] or "문서"


def _cell_to_str(v) -> str:
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v).strip()


def read_excel(path: str):
    """첫 행을 헤더로 읽어 (headers, rows[dict]) 반환. 빈 행은 건너뛴다."""
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    ws = wb.active
    rows_iter = ws.iter_rows(values_only=True)
    try:
        header_row = next(rows_iter)
    except StopIteration:
        wb.close()
        return [], []
    headers = [str(h).strip() for h in header_row if h is not None and str(h).strip()]
    n = len(headers)
    rows = []
    for raw in rows_iter:
        if raw is None or all(v is None or str(v).strip() == "" for v in raw[:n]):
            continue
        row = {}
        for i, h in enumerate(headers):
            v = raw[i] if i < len(raw) else None
            row[h] = _cell_to_str(v)
        rows.append(row)
    wb.close()
    return headers, rows


def render_template(text: str, row: dict) -> str:
    """'{{컬럼명}}' 자리표시를 행 데이터로 채운다."""
    return PLACEHOLDER.sub(lambda m: row.get(m.group(1).strip(), ""), text or "")


def find_unknown_placeholders(text: str, headers: list) -> list:
    """텍스트에 있지만 엑셀 헤더에 없는 자리표시 목록 (실행 전 오타 검증).

    {{컬럼}}(템플릿)과 {컬럼}(파일명 패턴) 두 형식을 모두 검사한다.
    """
    known = {h.strip() for h in headers}
    seen = []
    for m in re.finditer(r"\{\{?([^{}]+)\}?\}", text or ""):
        name = m.group(1).strip()
        if name and name not in known and name not in seen:
            seen.append(name)
    return seen


def fill_filename(pattern: str, row: dict) -> str:
    """'{이름}_성과평가서' 같은 파일명 패턴을 행 데이터로 채운다(확장자 제외)."""
    name = re.sub(r"\{([^{}]+)\}",
                  lambda m: row.get(m.group(1).strip(), ""),
                  pattern or "").strip()
    return sanitize_filename(name or "문서")


def unique_path(path: str) -> str:
    """동일 파일명 충돌 시 _2, _3 … 를 붙여 고유 경로를 만든다."""
    if not os.path.exists(path):
        return path
    base, ext = os.path.splitext(path)
    i = 2
    while os.path.exists(f"{base}_{i}{ext}"):
        i += 1
    return f"{base}_{i}{ext}"
